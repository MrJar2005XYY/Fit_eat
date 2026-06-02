from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# === Colors ===
PRIMARY = RGBColor(0x1A, 0x5C, 0xB0)
GREEN = RGBColor(0x34, 0xD3, 0x99)
DARK = RGBColor(0x1F, 0x29, 0x37)
GRAY = RGBColor(0x6B, 0x72, 0x80)
LIGHT_BG = RGBColor(0xF0, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AI_BG = RGBColor(0xE8, 0xF5, 0xE9)
ACCENT_ORANGE = RGBColor(0xFF, 0x98, 0x00)
LIGHT_GREEN = RGBColor(0xE8, 0xF5, 0xE9)
LIGHT_ORANGE = RGBColor(0xFF, 0xF3, 0xE0)
LIGHT_BLUE = RGBColor(0xE3, 0xF2, 0xFD)

SW = Inches(13.333)  # slide width
SH = Inches(7.5)     # slide height
MR = Inches(0.6)     # right margin
ML = Inches(0.6)     # left margin
CONTENT_W = Inches(12.133)  # SW - ML - MR

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH

def add_bg(slide, color=WHITE):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color

def add_rect(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape

def add_text(slide, left, top, width, height, text, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return tf

def add_bullets(slide, left, top, width, height, items, size=14, color=DARK):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(4)
    return tf

def add_header(slide, title, subtitle=""):
    add_rect(slide, Inches(0), Inches(0), SW, Inches(1.1), PRIMARY)
    add_text(slide, ML, Inches(0.15), Inches(10), Inches(0.5),
             title, size=30, bold=True, color=WHITE)
    if subtitle:
        add_text(slide, ML, Inches(0.65), Inches(10), Inches(0.35),
                 subtitle, size=14, color=RGBColor(0xBB, 0xCC, 0xEE))

def add_page_num(slide, num):
    add_text(slide, Inches(12.3), Inches(7.1), Inches(0.8), Inches(0.3),
             f"{num}/20", size=9, color=GRAY, align=PP_ALIGN.RIGHT)

def add_ai_box(slide, left, top, width, question, answers, adjusts=None):
    """Add an AI interaction block. Returns the height used."""
    line_h = Inches(0.22)
    q_lines = max(1, len(question) // 35 + 1)
    a_lines = len(answers)
    adj_lines = len(adjusts) if adjusts else 0

    h = Inches(0.35) + line_h * q_lines + Inches(0.15) + line_h * a_lines
    if adjusts:
        h += Inches(0.15) + line_h * adj_lines
    h = max(h, Inches(1.5))
    h = min(h, Inches(3.8))

    add_rect(slide, left, top, width, h, AI_BG)

    y = top + Inches(0.08)
    add_text(slide, left + Inches(0.15), y, width - Inches(0.3), Inches(0.25),
             "我们问 AI：", size=11, bold=True, color=PRIMARY)
    y += Inches(0.25)
    q_h = line_h * q_lines + Inches(0.05)
    add_text(slide, left + Inches(0.15), y, width - Inches(0.3), q_h,
             f'"{question}"', size=10, color=DARK)
    y += q_h + Inches(0.08)

    add_text(slide, left + Inches(0.15), y, width - Inches(0.3), Inches(0.22),
             "AI 回答：", size=11, bold=True, color=PRIMARY)
    y += Inches(0.22)
    for item in answers:
        add_text(slide, left + Inches(0.3), y, width - Inches(0.5), line_h,
                 f"• {item}", size=10, color=DARK)
        y += line_h

    if adjusts:
        y += Inches(0.08)
        add_text(slide, left + Inches(0.15), y, width - Inches(0.3), Inches(0.22),
                 "我们的调整：", size=11, bold=True, color=ACCENT_ORANGE)
        y += Inches(0.22)
        for item in adjusts:
            add_text(slide, left + Inches(0.3), y, width - Inches(0.5), line_h,
                     f"• {item}", size=10, color=DARK)
            y += line_h

    return h

# ============================================================
# SLIDE 1: Cover
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, PRIMARY)
add_text(slide, Inches(1), Inches(1.5), Inches(11), Inches(1),
         "轻食刻", size=56, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(2.6), Inches(11), Inches(0.5),
         "Fresh & Vitality", size=26, color=RGBColor(0xCC, 0xDD, 0xFF), align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(3.3), Inches(11), Inches(0.4),
         "AI 辅助开发的健康饮食管理应用", size=18, color=WHITE, align=PP_ALIGN.CENTER)
add_rect(slide, Inches(5.5), Inches(4.0), Inches(2.3), Inches(0.04), GREEN)
add_text(slide, Inches(1), Inches(4.5), Inches(11), Inches(0.4),
         "课程期末项目答辩", size=16, color=RGBColor(0xAA, 0xBB, 0xDD), align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(5.1), Inches(11), Inches(0.4),
         "团队成员：成员A · 成员B · 成员C · 成员D · 成员E", size=14, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(6.2), Inches(11), Inches(0.3),
         "2026年6月", size=13, color=RGBColor(0x99, 0xAA, 0xCC), align=PP_ALIGN.CENTER)
add_page_num(slide, 1)

# ============================================================
# SLIDE 2: TOC
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "目录", "CONTENTS")

sections = [
    ("01", "需求分析", "成员A", "竞品分析 · 用户画像 · 功能清单"),
    ("02", "原型系统设计", "成员B", "设计规范 · 移动端原型 · 组件库"),
    ("03", "前端开发", "成员C", "技术选型 · 页面开发 · 数据可视化"),
    ("04", "后端开发", "成员D", "架构设计 · 数据库 · API开发"),
    ("05", "测试与总结", "成员E", "测试过程 · 亮点 · 不足与展望"),
]

for i, (num, title, member, desc) in enumerate(sections):
    y = Inches(1.5) + Inches(1.1) * i
    add_rect(slide, ML, y, CONTENT_W, Inches(0.95), LIGHT_BG)
    add_text(slide, Inches(0.9), y + Inches(0.08), Inches(0.8), Inches(0.45),
             num, size=26, bold=True, color=PRIMARY)
    add_text(slide, Inches(1.9), y + Inches(0.08), Inches(4), Inches(0.4),
             title, size=18, bold=True, color=DARK)
    add_text(slide, Inches(1.9), y + Inches(0.5), Inches(6), Inches(0.3),
             desc, size=12, color=GRAY)
    add_text(slide, Inches(10.5), y + Inches(0.25), Inches(2), Inches(0.4),
             member, size=14, bold=True, color=PRIMARY, align=PP_ALIGN.RIGHT)
add_page_num(slide, 2)

# ============================================================
# SLIDE 3: A - 需求调研
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "一、需求分析", "成员A · 需求调研与 AI 辅助分析")

# Left column
add_text(slide, ML, Inches(1.3), Inches(5.5), Inches(0.35),
         "市场痛点", size=20, bold=True, color=DARK)
add_bullets(slide, ML, Inches(1.7), Inches(5.5), Inches(2.5), [
    "• 市面饮食 App 记录流程繁琐，用户留存率低",
    "• 缺乏个性化方案，千人一面的建议",
    "• 社交属性弱，用户缺少动力坚持",
    "• 界面设计复杂，学习成本高",
], size=14)

# Right column - AI block
add_ai_box(slide, Inches(6.5), Inches(1.3), Inches(6.2),
    "我想做一个健康饮食管理App，面向20-35岁年轻人，帮我分析市面上类似App的优缺点。",
    [
        "竞品分析：薄荷健康（记录繁琐）、MyFitnessPal（界面复杂）",
        "差异化建议：简化记录流程、AI个性化方案",
        "核心方向：社区互动激励 + 游戏化成就系统",
    ],
    [
        "选择「AI定制 + 社区互动」作为核心卖点",
        "去掉复杂热量数据库，改用简化记录",
    ])

# Bottom summary
add_text(slide, ML, Inches(5.2), CONTENT_W, Inches(0.35),
         "最终方向", size=18, bold=True, color=DARK)
add_rect(slide, ML, Inches(5.6), CONTENT_W, Inches(1.2), LIGHT_BG)
add_bullets(slide, Inches(0.9), Inches(5.7), Inches(11.5), Inches(1), [
    "• 轻食刻 = 简化饮食记录 + AI 个性化定制 + 社区互动激励",
    "• 核心价值：让健康饮食成为自然生活方式，通过社交和游戏化提升用户留存",
], size=14)
add_page_num(slide, 3)

# ============================================================
# SLIDE 4: A - 用户画像
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "一、需求分析", "成员A · 用户画像与使用场景")

personas = [
    ("小林 · 22岁 · 大学生", "减脂期", "想控制饮食但不知道怎么算热量，希望有简单易用的记录工具"),
    ("小张 · 28岁 · 上班族", "健康管理", "工作忙没时间研究营养学，希望有AI帮我规划每日饮食"),
    ("小王 · 32岁 · 健身爱好者", "增肌期", "需要精确追踪蛋白质摄入，希望有社区交流健身餐经验"),
]

for i, (name, scene, desc) in enumerate(personas):
    x = ML + Inches(3.95) * i
    add_rect(slide, x, Inches(1.3), Inches(3.7), Inches(2.2), LIGHT_BG)
    add_text(slide, x + Inches(0.2), Inches(1.4), Inches(3.3), Inches(0.3),
             name, size=15, bold=True, color=PRIMARY)
    add_rect(slide, x + Inches(0.2), Inches(1.8), Inches(1.1), Inches(0.3), GREEN)
    add_text(slide, x + Inches(0.25), Inches(1.82), Inches(1), Inches(0.25),
             scene, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.2), Inches(2.3), Inches(3.3), Inches(0.9),
             desc, size=12, color=DARK)

# AI block at bottom
add_ai_box(slide, ML, Inches(3.8), CONTENT_W,
    "根据用户画像，帮我写出3个典型的用户故事（User Story）。",
    [
        "作为大学生，我希望简单记录每餐食物，以便控制每日热量摄入",
        "作为上班族，我希望AI根据身体数据生成饮食方案，以便省去研究营养学的时间",
        "作为健身爱好者，我希望查看蛋白质摄入的详细图表，以便调整增肌饮食计划",
    ])
add_page_num(slide, 4)

# ============================================================
# SLIDE 5: A - 功能清单
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "一、需求分析", "成员A · 功能需求清单（AI 多轮对话细化）")

rounds = [
    ("第1轮：粗略功能列表", "AI 给出5大模块：饮食记录、营养追踪、AI定制、社区互动、个人中心"),
    ("第2轮：追问子功能", "我们问：社区功能具体包含哪些？AI 细化为：发帖、点赞、评论、关注、分享"),
    ("第3轮：API级别细化", "我们问：帮我设计具体的API端点。AI 产出7个蓝图、33个接口的详细设计"),
]

for i, (rtitle, desc) in enumerate(rounds):
    y = Inches(1.3) + Inches(0.85) * i
    add_rect(slide, ML, y, CONTENT_W, Inches(0.75), LIGHT_BG)
    add_text(slide, Inches(0.9), y + Inches(0.05), Inches(4), Inches(0.3),
             rtitle, size=14, bold=True, color=PRIMARY)
    add_text(slide, Inches(0.9), y + Inches(0.35), Inches(11.5), Inches(0.35),
             desc, size=13, color=DARK)

add_text(slide, ML, Inches(4.0), CONTENT_W, Inches(0.35),
         "最终确定的 5 大核心功能模块：", size=16, bold=True, color=DARK)

features = ["饮食记录\n拍照/手动记录", "营养追踪\n卡路里+宏量", "AI定制\n个性化方案", "社区互动\n分享+激励", "成就系统\n勋章+打卡"]
for i, feat in enumerate(features):
    x = ML + Inches(2.4) * i
    add_rect(slide, x, Inches(4.5), Inches(2.15), Inches(0.85), PRIMARY)
    add_text(slide, x + Inches(0.1), Inches(4.55), Inches(1.95), Inches(0.75),
             feat, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_page_num(slide, 5)

# ============================================================
# SLIDE 6: B - 设计规范
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "二、原型系统设计", "成员B · 设计规范制定")

# Left: AI block
add_ai_box(slide, ML, Inches(1.3), Inches(6.2),
    "我要做一个健康饮食App，风格是现代极简+软触觉元素，请帮我制定完整的设计规范。",
    [
        "主色 #006e1c（健康绿）、辅色 #8b5000（活力橙）",
        "字体 Plus Jakarta Sans / PingFang SC",
        "圆角 16px、毛玻璃阴影效果",
        "Material Design 3 色彩体系",
    ],
    [
        "调整主色为 #1A5CB0 蓝色（更专业）",
        "保留绿色作为强调色",
    ])

# Right: Color palette
add_text(slide, Inches(7.2), Inches(1.3), Inches(5.5), Inches(0.35),
         "色彩体系", size=16, bold=True, color=DARK)
colors = [
    ("Primary", PRIMARY), ("Accent", GREEN), ("Orange", ACCENT_ORANGE),
    ("Dark", DARK), ("Gray", GRAY), ("BG", LIGHT_BG),
]
for i, (name, rgb) in enumerate(colors):
    x = Inches(7.2) + Inches(0.9) * i
    add_rect(slide, x, Inches(1.75), Inches(0.7), Inches(0.65), rgb)
    add_text(slide, x, Inches(2.45), Inches(0.7), Inches(0.2),
             name, size=8, color=DARK, align=PP_ALIGN.CENTER)

# Typography
add_text(slide, Inches(7.2), Inches(2.85), Inches(5.5), Inches(0.3),
         "字体与规范", size=16, bold=True, color=DARK)
add_bullets(slide, Inches(7.2), Inches(3.2), Inches(5.5), Inches(2.5), [
    "  Headline: 30px / Bold",
    "  Body: 16px / Regular",
    "  Label: 14px / Medium",
    "  圆角: 16px (卡片) / 8px (按钮)",
    "  阴影: 0px 4px 20px rgba(0,110,28,0.06)",
], size=12)
add_page_num(slide, 6)

# ============================================================
# SLIDE 7: B - 移动端原型
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "二、原型系统设计", "成员B · 移动端原型设计")

# AI block top-left
add_ai_box(slide, ML, Inches(1.3), Inches(7),
    "帮我设计饮食记录App首页布局，包含卡路里环形图、今日餐单、营养进度条，用HTML+Tailwind实现。",
    [
        "HTML结构：顶部导航 → 卡路里环 → 餐单卡片 → 进度条",
        "Tailwind样式：flex布局、圆角卡片、渐变背景",
    ],
    [
        "调整卡片间距和阴影参数",
        "添加毛玻璃效果的顶部导航",
    ])

# Phone mockups right
pages = ["首页", "记录页", "社区页", "个人中心"]
for i, name in enumerate(pages):
    x = Inches(8.2) + Inches(1.25) * i
    add_rect(slide, x, Inches(1.3), Inches(1.1), Inches(2.5), LIGHT_BG)
    add_text(slide, x + Inches(0.05), Inches(2.2), Inches(1), Inches(0.3),
             "[截图]", size=9, color=GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.05), Inches(2.6), Inches(1), Inches(0.3),
             name, size=11, bold=True, color=DARK, align=PP_ALIGN.CENTER)

# Bottom: design points
add_text(slide, ML, Inches(4.3), CONTENT_W, Inches(0.3),
         "设计要点", size=16, bold=True, color=DARK)
add_bullets(slide, ML, Inches(4.7), CONTENT_W, Inches(2.5), [
    "• 卡路里环形图：SVG实现，支持动画绘制，直观展示摄入/剩余/消耗",
    "• 餐单卡片：横向滚动，早餐/午餐/晚餐/加餐四张卡片",
    "• 营养进度条：蛋白质/碳水/脂肪三条，颜色区分，动画填充",
    "• 顶部导航：毛玻璃效果，固定定位，头像+Logo+搜索",
], size=13)
add_page_num(slide, 7)

# ============================================================
# SLIDE 8: B - 桌面端 + 组件库
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "二、原型系统设计", "成员B · 桌面端适配与组件库")

# AI block left
add_ai_box(slide, ML, Inches(1.3), Inches(6.2),
    "移动端页面需要适配桌面端，帮我设计Bento Grid布局的桌面版方案。",
    [
        "Bento Grid：不规则网格布局，大卡片+小卡片组合",
        "顶部改为水平导航栏，左侧添加侧边栏",
        "内容区采用3列自适应布局",
    ],
    [
        "保留移动端底部导航作为备选",
        "桌面端社区页改为瀑布流3列布局",
    ])

# Component library right
add_text(slide, Inches(7.2), Inches(1.3), Inches(5.5), Inches(0.3),
         "核心组件库（AI 辅助生成）", size=16, bold=True, color=DARK)
components = [
    ("环形进度图", "SVG圆形图表\n自定义颜色+动画"),
    ("柱状图", "周热量趋势\n7天数据对比"),
    ("雷达图", "六维营养评估\n多边形叠加"),
    ("进度条", "宏量营养素\n蛋白质/碳水/脂肪"),
]
for i, (name, desc) in enumerate(components):
    x = Inches(7.2) + Inches(1.4) * i
    add_rect(slide, x, Inches(1.7), Inches(1.25), Inches(2.0), LIGHT_BG)
    add_text(slide, x + Inches(0.05), Inches(1.8), Inches(1.15), Inches(0.25),
             name, size=11, bold=True, color=PRIMARY, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.05), Inches(2.1), Inches(1.15), Inches(1.3),
             desc, size=9, color=DARK, align=PP_ALIGN.CENTER)

# AI block bottom
add_ai_box(slide, ML, Inches(4.2), CONTENT_W,
    "帮我用SVG实现一个环形进度图组件，支持自定义颜色和百分比动画。",
    [
        "使用SVG的circle元素，stroke-dasharray实现进度效果",
        "CSS transition实现平滑动画，支持传入颜色、百分比等参数",
    ])
add_page_num(slide, 8)

# ============================================================
# SLIDE 9: B - 交互组件
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "二、原型系统设计", "成员B · 交互组件与设计亮点")

add_text(slide, ML, Inches(1.3), CONTENT_W, Inches(0.3),
         "交互组件清单", size=18, bold=True, color=DARK)

interactions = [
    ("日期选择器", "横向滚动日期，左右滑动切换周"),
    ("分类筛选", "Chip标签，支持多选/单选切换"),
    ("Tab切换", "内容区域切换，带滑动指示器"),
    ("FAB按钮", "浮动操作按钮，快速添加记录"),
    ("Toast提示", "操作反馈，自动消失"),
    ("下拉刷新", "移动端下拉刷新数据"),
]

for i, (name, desc) in enumerate(interactions):
    col = i % 2
    row = i // 2
    x = ML + Inches(6.2) * col
    y = Inches(1.75) + Inches(0.7) * row
    add_rect(slide, x, y, Inches(5.9), Inches(0.6), LIGHT_BG)
    add_text(slide, x + Inches(0.15), y + Inches(0.05), Inches(1.6), Inches(0.25),
             name, size=13, bold=True, color=PRIMARY)
    add_text(slide, x + Inches(0.15), y + Inches(0.3), Inches(5.5), Inches(0.25),
             desc, size=12, color=GRAY)

add_text(slide, ML, Inches(4.0), CONTENT_W, Inches(0.3),
         "设计亮点总结", size=18, bold=True, color=DARK)
add_bullets(slide, ML, Inches(4.4), CONTENT_W, Inches(3), [
    "• 一致的设计语言：所有页面遵循统一的色彩、字体、圆角规范",
    "• 丰富的数据可视化：环形图、柱状图、雷达图、进度条4种图表",
    "• 响应式设计：Mobile First + 桌面端 Bento Grid 适配",
    "• 微交互动画：hover缩放、页面淡入、进度条动画",
], size=14)
add_page_num(slide, 9)

# ============================================================
# SLIDE 10: C - 前端技术选型
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "三、前端开发", "成员C · 技术选型与页面开发")

# AI block left
add_ai_box(slide, ML, Inches(1.3), Inches(6.5),
    "做一个轻量级健康饮食Web应用，不需要React/Vue，帮我推荐最简单的前端技术方案。",
    [
        "纯HTML + Tailwind CDN + Vanilla JS",
        "无需构建工具（webpack/vite），零配置",
        "Flask直接serve静态文件，前后端同域",
        "Google Material Symbols图标库（CDN）",
    ],
    [
        "确认采用此方案，开发效率最高",
        "添加common.css自定义全局样式",
        "封装api.js统一管理API调用",
    ])

# Tech cards right
add_text(slide, Inches(7.5), Inches(1.3), Inches(5.5), Inches(0.3),
         "前端技术栈", size=18, bold=True, color=DARK)
techs = [
    ("HTML5", "多页面架构"),
    ("Tailwind", "CDN引入"),
    ("Vanilla JS", "原生JS"),
    ("SVG", "数据可视化"),
]
for i, (name, desc) in enumerate(techs):
    x = Inches(7.5) + Inches(1.4) * i
    add_rect(slide, x, Inches(1.7), Inches(1.25), Inches(1.0), PRIMARY)
    add_text(slide, x + Inches(0.05), Inches(1.8), Inches(1.15), Inches(0.3),
             name, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.05), Inches(2.15), Inches(1.15), Inches(0.4),
             desc, size=10, color=RGBColor(0xCC, 0xDD, 0xFF), align=PP_ALIGN.CENTER)

# File structure
add_text(slide, ML, Inches(4.3), CONTENT_W, Inches(0.3),
         "前端文件结构", size=16, bold=True, color=DARK)
add_bullets(slide, ML, Inches(4.7), CONTENT_W, Inches(2.5), [
    "轻饮食软件前端/",
    "├── index.html          # 入口（自动跳转登录/首页）",
    "├── css/common.css      # 全局样式（玻璃态、阴影、导航）",
    "├── js/api.js           # API客户端（所有fetch调用）",
    "├── js/common.js        # 导航、认证检查、工具函数",
    "├── js/components.js    # 可复用组件（环形图、柱状图、雷达图）",
    "└── pages/              # 10个独立页面",
], size=12)
add_page_num(slide, 10)

# ============================================================
# SLIDE 11: C - 页面开发
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "三、前端开发", "成员C · 典型页面开发过程")

add_ai_box(slide, ML, Inches(1.3), CONTENT_W,
    "帮我用HTML+Tailwind实现首页设计稿，包含顶部导航、卡路里环形图、餐单卡片横向滚动、营养素进度条。",
    [
        "AI产出基础HTML结构 + Tailwind类名",
        "生成卡路里环形图的SVG代码",
        "实现餐单卡片的横向滚动flex布局",
    ],
    [
        "调整颜色变量匹配设计规范",
        "优化移动端触摸滚动体验",
        "添加数据绑定（从API获取真实数据）",
    ])

add_text(slide, ML, Inches(4.3), CONTENT_W, Inches(0.3),
         "AI 辅助编码的工作流", size=16, bold=True, color=DARK)

workflow = [
    ("1.描述需求", "给AI发送设计稿\n+功能描述"),
    ("2.AI生成", "产出HTML/CSS/JS\n基础代码"),
    ("3.我们调整", "修改样式\n适配规范"),
    ("4.数据对接", "连接api.js\n获取数据"),
    ("5.测试优化", "跨浏览器\n修复兼容"),
]
for i, (step, desc) in enumerate(workflow):
    x = ML + Inches(2.4) * i
    add_rect(slide, x, Inches(4.7), Inches(2.15), Inches(1.6), LIGHT_BG)
    add_text(slide, x + Inches(0.1), Inches(4.8), Inches(1.95), Inches(0.3),
             step, size=12, bold=True, color=PRIMARY, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.1), Inches(5.2), Inches(1.95), Inches(0.8),
             desc, size=11, color=DARK, align=PP_ALIGN.CENTER)
add_page_num(slide, 11)

# ============================================================
# SLIDE 12: C - 数据可视化
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "三、前端开发", "成员C · 数据可视化组件实现")

# AI block left
add_ai_box(slide, ML, Inches(1.3), Inches(6.5),
    "用原生JS+SVG实现六维雷达图组件，数据从API获取，支持动画绘制。",
    [
        "使用SVG polygon绘制六边形",
        "6维度：能量/蛋白质/脂肪/纤维/维生素/矿物质",
        "CSS @keyframes实现绘制动画",
        "封装为RadarChart类，支持动态更新",
    ],
    [
        "调整刻度线为3层（60%/80%/100%）",
        "添加维度标签自动避让",
    ])

# Chart cards right
add_text(slide, Inches(7.5), Inches(1.3), Inches(5.5), Inches(0.3),
         "4种图表组件", size=16, bold=True, color=DARK)
charts = [
    ("环形进度图", "卡路里追踪\nSVG circle"),
    ("柱状图", "周热量趋势\nSVG rect"),
    ("雷达图", "六维营养\nSVG polygon"),
    ("进度条", "宏量营养素\ndiv+CSS"),
]
for i, (name, desc) in enumerate(charts):
    x = Inches(7.5) + Inches(1.4) * i
    add_rect(slide, x, Inches(1.7), Inches(1.25), Inches(1.8), LIGHT_BG)
    add_text(slide, x + Inches(0.05), Inches(1.8), Inches(1.15), Inches(0.25),
             name, size=11, bold=True, color=PRIMARY, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.05), Inches(2.15), Inches(1.15), Inches(1.0),
             desc, size=10, color=DARK, align=PP_ALIGN.CENTER)

# Code snippet
add_text(slide, ML, Inches(4.6), CONTENT_W, Inches(0.3),
         "雷达图核心代码", size=14, bold=True, color=DARK)
add_rect(slide, ML, Inches(4.95), CONTENT_W, Inches(2.2), RGBColor(0x26, 0x2B, 0x35))
add_bullets(slide, Inches(0.9), Inches(5.05), Inches(11.5), Inches(2), [
    "// 计算六边形顶点坐标",
    "const points = data.map((val, i) => {",
    "  const angle = (Math.PI / 3) * i - Math.PI / 2;",
    "  const r = (val / max) * radius;",
    "  return `${cx + r*Math.cos(angle)},${cy + r*Math.sin(angle)}`;",
    "}).join(' ');",
    "polygon.setAttribute('points', points);",
], size=11, color=RGBColor(0xA6, 0xE2, 0x2E))
add_page_num(slide, 12)

# ============================================================
# SLIDE 13: C - 交互动画
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "三、前端开发", "成员C · 交互效果与动画实现")

# AI block top
add_ai_box(slide, ML, Inches(1.3), CONTENT_W,
    "帮我实现底部导航栏的Tab切换效果，当前页高亮，页面切换有淡入动画。",
    [
        "5个Tab：首页/发现/记录/社区/我的",
        "当前Tab图标填充+文字变色",
        "页面切换使用CSS fadeIn动画",
    ],
    [
        "添加Tab切换的微交互缩放效果",
        "适配桌面端改为顶部水平导航",
    ])

# Interaction cards
add_text(slide, ML, Inches(4.2), CONTENT_W, Inches(0.3),
         "AI 辅助实现的交互效果", size=16, bold=True, color=DARK)

effects = [
    ("底部导航", "Tab高亮+图标填充\n页面淡入动画"),
    ("卡片微交互", "hover:scale(1.02)\ncubic-bezier缓动"),
    ("毛玻璃背景", "backdrop-filter:blur\n半透明白色背景"),
    ("入场动画", "fadeInUp动画\n滚动触发Observer"),
]
for i, (name, desc) in enumerate(effects):
    x = ML + Inches(3.05) * i
    add_rect(slide, x, Inches(4.6), Inches(2.8), Inches(1.4), LIGHT_BG)
    add_text(slide, x + Inches(0.1), Inches(4.7), Inches(2.6), Inches(0.25),
             name, size=13, bold=True, color=PRIMARY, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.1), Inches(5.0), Inches(2.6), Inches(0.8),
             desc, size=11, color=DARK, align=PP_ALIGN.CENTER)

# Efficiency
add_text(slide, ML, Inches(6.2), CONTENT_W, Inches(0.3),
         "前端开发效率：AI 辅助平均节省约 60% 编码时间", size=14, bold=True, color=GREEN)
add_page_num(slide, 13)

# ============================================================
# SLIDE 14: D - 后端架构
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "四、后端开发", "成员D · 架构设计与技术选型")

# AI block left
add_ai_box(slide, ML, Inches(1.3), Inches(6.5),
    "用Python Flask做饮食管理应用的后端，帮我设计项目结构、目录组织和蓝图划分。",
    [
        "工厂模式：create_app()函数初始化",
        "models/目录放SQLAlchemy数据模型",
        "routes/目录按功能分蓝图",
        "admin/目录放Flask-Admin管理后台",
    ],
    [
        "选择Flask而非Django（轻量、快速原型）",
        "选择SQLite而非MySQL（零配置）",
    ])

# Architecture diagram right
add_text(slide, Inches(7.5), Inches(1.3), Inches(5.5), Inches(0.3),
         "系统架构图", size=18, bold=True, color=DARK)
layers = [
    ("浏览器（前端）", "HTML + Tailwind CSS + JS", LIGHT_BLUE),
    ("Flask 服务器", "7个蓝图 + 静态文件 + Admin", LIGHT_GREEN),
    ("SQLite 数据库", "11张表 + ORM映射", LIGHT_ORANGE),
]
for i, (name, desc, bg) in enumerate(layers):
    y = Inches(1.75) + Inches(1.15) * i
    add_rect(slide, Inches(7.5), y, Inches(5.5), Inches(0.95), bg)
    add_text(slide, Inches(7.7), y + Inches(0.08), Inches(5), Inches(0.3),
             name, size=15, bold=True, color=DARK)
    add_text(slide, Inches(7.7), y + Inches(0.42), Inches(5), Inches(0.3),
             desc, size=12, color=GRAY)
    if i < 2:
        add_text(slide, Inches(9.8), y + Inches(0.95), Inches(1), Inches(0.25),
                 "↕", size=16, color=GRAY, align=PP_ALIGN.CENTER)

# Tech choices
add_text(slide, ML, Inches(5.2), CONTENT_W, Inches(0.3),
         "技术选型理由", size=16, bold=True, color=DARK)
choices = [
    ("Flask", "轻量灵活，适合快速原型，学习成本低"),
    ("SQLAlchemy", "Python ORM，模型定义简洁，迁移方便"),
    ("Flask-Login", "轻量认证方案，session管理简单"),
    ("Flask-Admin", "开箱即用的管理后台，省去自建成本"),
]
for i, (tech, reason) in enumerate(choices):
    y = Inches(5.6) + Inches(0.35) * i
    add_text(slide, ML, y, Inches(1.8), Inches(0.3),
             tech, size=12, bold=True, color=PRIMARY)
    add_text(slide, Inches(2.5), y, Inches(10), Inches(0.3),
             reason, size=12, color=DARK)
add_page_num(slide, 14)

# ============================================================
# SLIDE 15: D - 数据库
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "四、后端开发", "成员D · 数据库设计")

# AI block left
add_ai_box(slide, ML, Inches(1.3), Inches(6.5),
    "帮我设计饮食管理应用的数据库，包含用户、食物、饮食记录、社区帖子、评论、点赞、关注、成就等表。",
    [
        "11张表覆盖用户/食物/记录/社区/成就",
        "外键关系：user→diet_record, post→comment",
        "联合唯一约束：likes表(user_id, post_id)",
    ],
    [
        "添加account_id字段（区别于登录ID）",
        "AI身体数据单独建表存储",
    ])

# Table list right
add_text(slide, Inches(7.5), Inches(1.3), Inches(5.5), Inches(0.3),
         "11 张数据表", size=16, bold=True, color=DARK)
tables = [
    ("users", "用户信息"),
    ("foods", "食物库"),
    ("diet_records", "饮食记录"),
    ("water_records", "饮水记录"),
    ("community_posts", "社区帖子"),
    ("comments", "评论"),
    ("likes", "点赞"),
    ("follows", "关注"),
    ("achievements", "成就定义"),
    ("user_achievements", "用户成就"),
    ("ai_body_data", "AI身体数据"),
]
for i, (name, desc) in enumerate(tables):
    y = Inches(1.7) + Inches(0.42) * i
    bg = LIGHT_BG if i % 2 == 0 else WHITE
    add_rect(slide, Inches(7.5), y, Inches(5.5), Inches(0.36), bg)
    add_text(slide, Inches(7.7), y + Inches(0.04), Inches(2), Inches(0.28),
             name, size=10, bold=True, color=PRIMARY)
    add_text(slide, Inches(9.8), y + Inches(0.04), Inches(3), Inches(0.28),
             desc, size=10, color=DARK)
add_page_num(slide, 15)

# ============================================================
# SLIDE 16: D - API
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "四、后端开发", "成员D · API 开发")

# AI block
add_ai_box(slide, ML, Inches(1.3), CONTENT_W,
    "帮我用Flask-Login实现用户登录注册API，密码用Werkzeug哈希，返回一个简单的token。",
    [
        "注册：generate_password_hash加密存储",
        "登录：check_password_hash验证，创建session",
        "返回user_{id}作为token（非JWT，简化实现）",
    ],
    [
        "token格式简化为user_{id}字符串",
        "前端localStorage存储token",
    ])

# API table
add_text(slide, ML, Inches(4.2), CONTENT_W, Inches(0.3),
         "7 个 API 蓝图 · 33 个接口", size=16, bold=True, color=DARK)

api_data = [
    ("auth", "/api/auth", "注册/登录/退出", "3"),
    ("user", "/api/user", "资料/账号/密码", "4"),
    ("diet", "/api/diet", "记录/卡路里/水/宏量/雷达", "9"),
    ("food", "/api/food", "列表/详情/搜索/收藏", "4"),
    ("community", "/api/community", "帖子/点赞/评论/关注", "7"),
    ("achievement", "/api/achievement", "成就/AI方案", "5"),
    ("upload", "/api/upload", "图片上传", "1"),
]

# Header row
for j, h in enumerate(["蓝图", "前缀", "功能", "接口数"]):
    x = ML + Inches(1.7) * j + (Inches(3) if j == 2 else 0)
    w = Inches(3) if j == 2 else Inches(1.5)
    add_text(slide, x, Inches(4.55), w, Inches(0.25),
             h, size=11, bold=True, color=PRIMARY)

for i, (name, prefix, func, count) in enumerate(api_data):
    y = Inches(4.85) + Inches(0.33) * i
    bg = LIGHT_BG if i % 2 == 0 else WHITE
    add_rect(slide, ML, y, CONTENT_W, Inches(0.28), bg)
    add_text(slide, Inches(0.8), y + Inches(0.02), Inches(1.5), Inches(0.24),
             name, size=10, bold=True, color=PRIMARY)
    add_text(slide, Inches(2.5), y + Inches(0.02), Inches(2), Inches(0.24),
             prefix, size=10, color=GRAY)
    add_text(slide, Inches(5), y + Inches(0.02), Inches(5.5), Inches(0.24),
             func, size=10, color=DARK)
    add_text(slide, Inches(11.5), y + Inches(0.02), Inches(1), Inches(0.24),
             count, size=10, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
add_page_num(slide, 16)

# ============================================================
# SLIDE 17: D - 数据导入
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "四、后端开发", "成员D · 数据导入与管理后台")

# AI block left
add_ai_box(slide, ML, Inches(1.3), Inches(6.5),
    "帮我写Python脚本，解析HowToCook开源项目的Markdown食谱文件，导入到SQLite数据库。",
    [
        "解析Markdown标题获取菜名",
        "正则匹配食材列表和用量",
        "批量插入Food表，自动去重",
    ],
    [
        "添加图片下载脚本（import_images.py）",
        "处理GitHub LFS图片链接",
    ])

# Admin panel right
add_text(slide, Inches(7.5), Inches(1.3), Inches(5.5), Inches(0.3),
         "Flask-Admin 管理后台", size=16, bold=True, color=DARK)
add_bullets(slide, Inches(7.5), Inches(1.7), Inches(5.5), Inches(2.5), [
    "  9个数据模型视图",
    "  支持搜索、过滤、排序",
    "  支持数据导出（CSV/Excel）",
    "  支持关联数据内联编辑",
    "  访问：localhost:5000/admin",
], size=13)

# AI block bottom
add_ai_box(slide, ML, Inches(4.5), CONTENT_W,
    "帮我配置Flask-Admin，给所有模型添加搜索、过滤和导出功能。",
    [
        "继承ModelView自定义每个视图",
        "column_searchable_list配置搜索字段",
        "can_export = True开启CSV导出",
    ])
add_page_num(slide, 17)

# ============================================================
# SLIDE 18: E - 测试
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "五、测试与总结", "成员E · 测试过程")

# AI block left
add_ai_box(slide, ML, Inches(1.3), Inches(6.5),
    "帮我列出这个饮食管理应用的测试用例，包括正常流程和边界情况。",
    [
        "登录：正确密码/错误密码/空输入",
        "饮食记录：正常添加/超大数值/负数",
        "社区：发帖/删帖/重复点赞",
    ],
    [
        "按模块整理为测试清单",
        "逐项手动执行并记录结果",
    ])

# Test results right
add_text(slide, Inches(7.5), Inches(1.3), Inches(5.5), Inches(0.3),
         "测试结果", size=16, bold=True, color=DARK)
test_results = [
    ("用户认证", "通过", "登录/注册/退出正常"),
    ("饮食记录", "通过", "增删改查+卡路里计算"),
    ("数据可视化", "通过", "4种图表渲染正常"),
    ("社区功能", "通过", "发帖/点赞/评论正常"),
    ("AI定制", "通过", "数据录入+方案生成"),
    ("成就系统", "通过", "条件判断+勋章解锁"),
    ("图片上传", "已知问题", "大图片偶尔超时"),
    ("响应式", "通过", "移动端/桌面端正常"),
]
for i, (module, status, desc) in enumerate(test_results):
    y = Inches(1.7) + Inches(0.48) * i
    bg = LIGHT_BG if i % 2 == 0 else WHITE
    add_rect(slide, Inches(7.5), y, Inches(5.5), Inches(0.4), bg)
    add_text(slide, Inches(7.7), y + Inches(0.05), Inches(1.3), Inches(0.28),
             module, size=11, bold=True, color=DARK)
    color = GREEN if status == "通过" else ACCENT_ORANGE
    add_text(slide, Inches(9.0), y + Inches(0.05), Inches(1.2), Inches(0.28),
             status, size=11, bold=True, color=color)
    add_text(slide, Inches(10.3), y + Inches(0.05), Inches(2.5), Inches(0.28),
             desc, size=10, color=GRAY)

# Bugs
add_text(slide, ML, Inches(5.3), CONTENT_W, Inches(0.3),
         "发现并修复的 Bug", size=14, bold=True, color=DARK)
add_bullets(slide, ML, Inches(5.7), CONTENT_W, Inches(1.5), [
    "• Bug 1：重复点赞不报错 → 添加唯一约束检查",
    "• Bug 2：编辑资料后头像不更新 → 修复缓存问题",
    "• Bug 3：食谱数据显示Markdown符号 → 添加格式清理",
], size=12)
add_page_num(slide, 18)

# ============================================================
# SLIDE 19: E - 亮点与不足
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_header(slide, "五、测试与总结", "成员E · 项目亮点与不足")

# Highlights left
add_text(slide, ML, Inches(1.3), Inches(5.5), Inches(0.3),
         "项目亮点", size=20, bold=True, color=GREEN)
highlights = [
    ("AI 深度参与", "全流程AI辅助，从需求到测试"),
    ("设计系统", "Material Design 3，统一规范"),
    ("数据可视化", "4种SVG图表，动画交互"),
    ("社区互动", "发帖/点赞/评论/关注"),
    ("成就系统", "游戏化勋章墙，激励打卡"),
    ("响应式", "移动端+桌面端双端适配"),
]
for i, (name, desc) in enumerate(highlights):
    col = i % 2
    row = i // 2
    x = ML + Inches(2.85) * col
    y = Inches(1.7) + Inches(0.6) * row
    add_rect(slide, x, y, Inches(2.65), Inches(0.5), LIGHT_GREEN)
    add_text(slide, x + Inches(0.1), y + Inches(0.03), Inches(2.45), Inches(0.22),
             name, size=12, bold=True, color=GREEN)
    add_text(slide, x + Inches(0.1), y + Inches(0.25), Inches(2.45), Inches(0.22),
             desc, size=10, color=DARK)

# Shortcomings right
add_text(slide, Inches(7.5), Inches(1.3), Inches(5.5), Inches(0.3),
         "不足与改进方向", size=20, bold=True, color=ACCENT_ORANGE)
shortcomings = [
    ("无自动化测试", "引入 pytest 测试框架"),
    ("AI用模拟数据", "接入真实 AI API"),
    ("图片存储简陋", "迁移至云存储（OSS）"),
    ("前后端未分离", "重构为 Vue/React SPA"),
]
for i, (name, plan) in enumerate(shortcomings):
    y = Inches(1.7) + Inches(0.6) * i
    add_rect(slide, Inches(7.5), y, Inches(5.5), Inches(0.5), LIGHT_ORANGE)
    add_text(slide, Inches(7.7), y + Inches(0.03), Inches(2.5), Inches(0.22),
             f"⚠ {name}", size=12, bold=True, color=ACCENT_ORANGE)
    add_text(slide, Inches(7.7), y + Inches(0.25), Inches(5), Inches(0.22),
             f"→ {plan}", size=10, color=DARK)

# AI efficiency summary
add_text(slide, ML, Inches(4.2), CONTENT_W, Inches(0.3),
         "AI 辅助开发效率总结", size=16, bold=True, color=DARK)
efficiency = [
    ("需求分析", "AI竞品分析+用户画像", "节省约50%"),
    ("原型设计", "AI生成规范+页面布局", "节省约60%"),
    ("前端开发", "AI生成HTML/CSS/JS", "节省约60%"),
    ("后端开发", "AI生成Flask架构+API", "节省约55%"),
    ("测试", "AI生成测试用例清单", "节省约40%"),
]
for i, (module, desc, saving) in enumerate(efficiency):
    y = Inches(4.6) + Inches(0.42) * i
    bg = LIGHT_BG if i % 2 == 0 else WHITE
    add_rect(slide, ML, y, CONTENT_W, Inches(0.36), bg)
    add_text(slide, Inches(0.8), y + Inches(0.04), Inches(1.5), Inches(0.28),
             module, size=11, bold=True, color=PRIMARY)
    add_text(slide, Inches(2.5), y + Inches(0.04), Inches(6), Inches(0.28),
             desc, size=11, color=DARK)
    add_text(slide, Inches(10), y + Inches(0.04), Inches(2.5), Inches(0.28),
             saving, size=11, bold=True, color=GREEN, align=PP_ALIGN.RIGHT)
add_page_num(slide, 19)

# ============================================================
# SLIDE 20: Thanks
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, PRIMARY)
add_text(slide, Inches(1), Inches(2), Inches(11), Inches(0.9),
         "感谢聆听", size=50, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(3.0), Inches(11), Inches(0.5),
         "Thank You", size=26, color=RGBColor(0xCC, 0xDD, 0xFF), align=PP_ALIGN.CENTER)
add_rect(slide, Inches(5.5), Inches(3.8), Inches(2.3), Inches(0.04), GREEN)
add_text(slide, Inches(1), Inches(4.3), Inches(11), Inches(0.4),
         "轻食刻 · Fresh & Vitality", size=18, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(5.0), Inches(11), Inches(0.35),
         "感谢老师的指导与帮助", size=14, color=RGBColor(0xAA, 0xBB, 0xDD), align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(5.6), Inches(11), Inches(0.5),
         "Q & A", size=30, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
add_page_num(slide, 20)

# === Save ===
out = "D:/健康饮食软件/轻食刻-期末答辩PPT.pptx"
prs.save(out)
print(f"Done: {out}")
print(f"Slides: {len(prs.slides)}")
